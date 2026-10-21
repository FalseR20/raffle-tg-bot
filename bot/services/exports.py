"""CSV, XLSX and print-ticket PDF exports."""

from __future__ import annotations

import csv
import io
import os
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from bot.db.enums import ApplicationStatus
from bot.db.models import Application, ParticipantTotal, Raffle
from bot.db.repositories import applications as apps_repo
from bot.domain.money import format_money_export
from bot.domain.tickets import compress_ranges
from bot.services.errors import ExportError
from bot.utils.csv import safe_csv_value, safe_excel_value
from bot.utils.time import format_datetime

APPLICATION_COLUMNS = [
    "raffle",
    "application_id",
    "payment_code",
    "organization",
    "payment_method",
    "user_id",
    "username",
    "first_name",
    "last_name",
    "amount_byn",
    "tickets_awarded",
    "tickets_total_after",
    "first_ticket",
    "last_ticket",
    "voided_first_ticket",
    "voided_last_ticket",
    "status",
    "submitted_at",
    "confirmed_at",
    "confirmed_by",
    "rejection_reason",
]

PARTICIPANT_COLUMNS = [
    "raffle",
    "user_id",
    "username",
    "first_name",
    "last_name",
    "confirmed_amount_byn",
    "tickets_issued",
    "ticket_ranges",
    "applications_count",
    "organizations",
]

TICKET_COLUMNS = [
    "ticket_number",
    "user_id",
    "username",
    "source_application_id",
    "organization",
    "payment_method",
    "amount_byn",
]

SUMMARY_COLUMNS = [
    "organization",
    "applications_count",
    "confirmed_count",
    "total_amount_byn",
]


@dataclass(frozen=True)
class ExportFile:
    data: bytes
    filename: str
    mime: str


@dataclass(frozen=True)
class ExportRequest:
    include_all_statuses: bool = False
    allowed_organization_ids: frozenset[int] | None = None
    is_org_rep: bool = False


@dataclass
class ExportData:
    applications: list[Application]
    participants: list[ParticipantTotal]
    all_applications: list[Application]
    ranges_by_user: dict[int, list[tuple[int, int]]]
    organizations_by_user: dict[int, set[str]]
    applications_by_user: dict[int, int]


async def _load_data(session: AsyncSession, raffle: Raffle, request: ExportRequest) -> ExportData:
    filters = apps_repo.ApplicationFilters(
        raffle_id=raffle.id,
        statuses=None if request.include_all_statuses else (ApplicationStatus.CONFIRMED,),
    )
    applications = await apps_repo.list_filtered(session, filters, limit=100_000)
    if request.allowed_organization_ids is not None:
        applications = [
            app for app in applications if app.organization_id in request.allowed_organization_ids
        ]

    all_filters = apps_repo.ApplicationFilters(raffle_id=raffle.id)
    all_applications = await apps_repo.list_filtered(session, all_filters, limit=100_000)
    if request.allowed_organization_ids is not None:
        all_applications = [
            app
            for app in all_applications
            if app.organization_id in request.allowed_organization_ids
        ]

    participants_stmt = select(ParticipantTotal).where(ParticipantTotal.raffle_id == raffle.id)
    participants = list((await session.execute(participants_stmt)).scalars())

    ranges_by_user: dict[int, list[tuple[int, int]]] = defaultdict(list)
    organizations_by_user: dict[int, set[str]] = defaultdict(set)
    applications_by_user: dict[int, int] = defaultdict(int)
    for app in all_applications:
        applications_by_user[app.user_id] += 1
        organizations_by_user[app.user_id].add(app.organization.name)
        if (
            app.status is ApplicationStatus.CONFIRMED
            and app.first_ticket_number is not None
            and app.last_ticket_number is not None
        ):
            ranges_by_user[app.user_id].append((app.first_ticket_number, app.last_ticket_number))

    return ExportData(
        applications=applications,
        participants=participants,
        all_applications=all_applications,
        ranges_by_user=ranges_by_user,
        organizations_by_user=organizations_by_user,
        applications_by_user=applications_by_user,
    )


def _application_row(app: Application, tz: ZoneInfo, raffle_title: str) -> list:
    user = app.user
    return [
        raffle_title,
        app.public_id,
        app.payment_code,
        app.organization.name,
        app.payment_method_title or "",
        app.user_id,
        user.username or "",
        user.first_name,
        user.last_name or "",
        format_money_export(app.amount) if app.amount is not None else "",
        app.tickets_awarded,
        app.tickets_total_after if app.tickets_total_after is not None else "",
        app.first_ticket_number if app.first_ticket_number is not None else "",
        app.last_ticket_number if app.last_ticket_number is not None else "",
        app.voided_first_ticket_number if app.voided_first_ticket_number is not None else "",
        app.voided_last_ticket_number if app.voided_last_ticket_number is not None else "",
        app.status.value,
        format_datetime(app.submitted_at, tz) if app.submitted_at else "",
        format_datetime(app.reviewed_at, tz) if app.reviewed_at else "",
        app.reviewed_by or "",
        app.rejection_reason or "",
    ]


def _participant_rows(data: ExportData, tz: ZoneInfo, raffle_title: str) -> list[list]:
    rows = []
    for total in sorted(data.participants, key=lambda item: item.user_id):
        user = total.user
        rows.append(
            [
                raffle_title,
                total.user_id,
                user.username or "",
                user.first_name,
                user.last_name or "",
                format_money_export(total.confirmed_amount),
                total.tickets_issued,
                compress_ranges(data.ranges_by_user.get(total.user_id, [])),
                data.applications_by_user.get(total.user_id, 0),
                ", ".join(sorted(data.organizations_by_user.get(total.user_id, set()))),
            ]
        )
    return rows


def _ticket_rows(data: ExportData, raffle_title: str) -> list[list]:
    rows = []
    for app in data.all_applications:
        if (
            app.status is not ApplicationStatus.CONFIRMED
            or app.first_ticket_number is None
            or app.last_ticket_number is None
        ):
            continue
        for number in range(app.first_ticket_number, app.last_ticket_number + 1):
            rows.append(
                [
                    number,
                    app.user_id,
                    app.user.username or "",
                    app.public_id,
                    app.organization.name,
                    app.payment_method_title or "",
                    format_money_export(app.amount) if app.amount is not None else "",
                ]
            )
    rows.sort(key=lambda row: row[0])
    return rows


def _organization_summary(data: ExportData) -> list[list]:
    stats: dict[str, list] = {}
    for app in data.all_applications:
        name = app.organization.name
        entry = stats.setdefault(name, [name, 0, 0, Decimal("0.00")])
        entry[1] += 1
        if app.status is ApplicationStatus.CONFIRMED:
            entry[2] += 1
            entry[3] += app.amount or Decimal("0.00")
    return [
        [name, counts, confirmed, format_money_export(amount)]
        for name, counts, confirmed, amount in sorted(stats.values())
    ]


def _summary_rows(data: ExportData, raffle: Raffle, *, include_organizations: bool) -> list[list]:
    rows: list[list] = []
    if include_organizations:
        rows.append(["organization", "applications_count", "confirmed_count", "total_amount_byn"])
        rows.extend(_organization_summary(data))
        rows.append([])

    confirmed = [app for app in data.all_applications if app.status is ApplicationStatus.CONFIRMED]
    confirmed_amount = sum(
        (app.amount or Decimal("0.00") for app in confirmed),
        Decimal("0.00"),
    )
    participants_count = len({app.user_id for app in data.all_applications})
    tickets_issued = sum(app.tickets_awarded for app in confirmed)

    rows.append(
        [
            "raffle",
            "prizes_count",
            "participants_count",
            "confirmed_amount_byn",
            "tickets_issued",
        ]
    )
    rows.append(
        [
            raffle.title,
            raffle.prizes_count,
            participants_count,
            format_money_export(confirmed_amount),
            tickets_issued,
        ]
    )
    return rows


def export_filename(raffle: Raffle, extension: str) -> str:
    suffix = datetime.now().strftime("%Y-%m-%d")
    return f"raffle_{raffle.slug}_{suffix}.{extension}"


def build_csv(data: ExportData, raffle: Raffle, tz: ZoneInfo) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow(APPLICATION_COLUMNS)
    for app in data.applications:
        row = _application_row(app, tz, raffle.title)
        writer.writerow([safe_csv_value(value) for value in row])
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def build_xlsx(data: ExportData, raffle: Raffle, tz: ZoneInfo, *, is_org_rep: bool) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    def add_sheet(
        title: str,
        headers: list[str],
        rows: list[list],
        *,
        prepend_headers: bool = True,
        money_columns: frozenset[int] = frozenset(),
    ) -> None:
        sheet = workbook.create_sheet(title)
        if prepend_headers:
            sheet.append(headers)
        for row in rows:
            converted = []
            for index, value in enumerate(row, start=1):
                if index in money_columns and value not in (None, ""):
                    try:
                        converted.append(Decimal(str(value)))
                        continue
                    except InvalidOperation:
                        pass
                converted.append(safe_excel_value(value))
            sheet.append(converted)
        header_font = Font(bold=True, color="FFFFFF")
        fill = PatternFill("solid", fgColor="305496")
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        if prepend_headers:
            sheet.auto_filter.ref = sheet.dimensions
        for row in sheet.iter_rows(min_row=2):
            for index in money_columns:
                if index <= len(row):
                    row[index - 1].number_format = "0.00"
        for index, header in enumerate(headers, start=1):
            width = max(
                12,
                min(
                    40,
                    max(
                        len(str(header)) + 2,
                        *(
                            [len(str(row[index - 1])) + 2 for row in rows if len(row) >= index]
                            or [0]
                        ),
                    ),
                ),
            )
            sheet.column_dimensions[get_column_letter(index)].width = width

    add_sheet(
        "Applications",
        APPLICATION_COLUMNS,
        [_application_row(app, tz, raffle.title) for app in data.applications],
        money_columns=frozenset({10}),
    )
    if not is_org_rep:
        add_sheet(
            "Participants",
            PARTICIPANT_COLUMNS,
            _participant_rows(data, tz, raffle.title),
            money_columns=frozenset({6}),
        )
    add_sheet(
        "Tickets",
        TICKET_COLUMNS,
        _ticket_rows(data, raffle.title),
        money_columns=frozenset({7}),
    )
    add_sheet(
        "Summary",
        SUMMARY_COLUMNS,
        _summary_rows(data, raffle, include_organizations=not is_org_rep),
        prepend_headers=False,
        money_columns=frozenset({4}),
    )

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


async def build_export(
    session: AsyncSession,
    *,
    raffle: Raffle,
    format_name: str,
    request: ExportRequest,
    tz: ZoneInfo,
) -> ExportFile:
    data = await _load_data(session, raffle, request)
    if format_name == "csv":
        return ExportFile(
            data=build_csv(data, raffle, tz),
            filename=export_filename(raffle, "csv"),
            mime="text/csv",
        )
    if format_name == "xlsx":
        return ExportFile(
            data=build_xlsx(data, raffle, tz, is_org_rep=request.is_org_rep),
            filename=export_filename(raffle, "xlsx"),
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    raise ExportError("Неизвестный формат экспорта.")


# ----------------------------- print tickets PDF -----------------------------

_FONT_NAME = "RaffleDejaVu"
_FONT_REGISTERED = False

_FONT_CANDIDATES = (
    os.environ.get("DEJAVU_FONT_PATH", ""),
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/TTF/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
)


def _register_font() -> str:
    global _FONT_REGISTERED
    if _FONT_REGISTERED:
        return _FONT_NAME
    for candidate in _FONT_CANDIDATES:
        if candidate and Path(candidate).is_file():
            pdfmetrics.registerFont(TTFont(_FONT_NAME, candidate))
            _FONT_REGISTERED = True
            return _FONT_NAME
    raise ExportError(
        "Не найден шрифт DejaVu Sans для PDF с кириллицей. "
        "Установите fonts-dejavu-core или задайте DEJAVU_FONT_PATH."
    )


def build_tickets_pdf(raffle: Raffle, *, tickets_per_page: int = 10) -> ExportFile:
    if raffle.status.value not in {"finished", "archived"}:
        raise ExportError("PDF с билетами доступен только после завершения розыгрыша.")
    if raffle.last_ticket_number < 1:
        raise ExportError("В розыгрыше нет выпущенных билетов.")
    font_name = _register_font()

    page_width, page_height = A4
    margin = 12 * mm
    columns = 2
    rows_per_column = max(1, tickets_per_page // columns)
    cell_width = (page_width - 2 * margin) / columns
    cell_height = (page_height - 2 * margin) / rows_per_column
    per_page = columns * rows_per_column

    output = io.BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4)
    pdf.setTitle(f"Билеты розыгрыша {raffle.title}")

    for index in range(raffle.last_ticket_number):
        position = index % per_page
        if position == 0:
            if index:
                pdf.showPage()
            pdf.setFont(font_name, 10)
            pdf.setFillColorRGB(0.35, 0.35, 0.35)
            pdf.drawString(
                margin,
                page_height - margin + 2 * mm,
                f"{raffle.title} · билеты {index + 1}–"
                f"{min(index + per_page, raffle.last_ticket_number)}",
            )
            pdf.setFillColorRGB(0, 0, 0)

        column = position % columns
        row = position // columns
        x = margin + column * cell_width
        y = page_height - margin - (row + 1) * cell_height

        pdf.setDash(2, 2)
        pdf.setStrokeColorRGB(0.65, 0.65, 0.65)
        pdf.rect(x, y, cell_width, cell_height)
        pdf.setDash()

        pdf.setFont(font_name, 34)
        pdf.drawCentredString(x + cell_width / 2, y + cell_height / 2 - 4 * mm, str(index + 1))
        pdf.setFont(font_name, 9)
        pdf.setFillColorRGB(0.3, 0.3, 0.3)
        pdf.drawCentredString(x + cell_width / 2, y + cell_height / 2 + 12 * mm, raffle.title[:48])
        pdf.setFont(font_name, 7)
        pdf.drawCentredString(
            x + cell_width / 2, y + 6 * mm, "Билет розыгрыша · сохраняйте до результатов"
        )
        pdf.setFillColorRGB(0, 0, 0)

    pdf.showPage()
    pdf.save()
    return ExportFile(
        data=output.getvalue(),
        filename=export_filename(raffle, "pdf"),
        mime="application/pdf",
    )


async def build_participants_export_rows(
    session: AsyncSession, raffle: Raffle, tz: ZoneInfo
) -> list[list]:
    """Helper used by tests and statistics; returns participant rows."""
    data = await _load_data(session, raffle, ExportRequest())
    return _participant_rows(data, tz, raffle.title)


__all__ = [
    "ExportData",
    "ExportFile",
    "ExportRequest",
    "build_csv",
    "build_export",
    "build_tickets_pdf",
    "build_xlsx",
    "export_filename",
]

"""Export, PDF and winner selection integration tests."""

from __future__ import annotations

import io
import random
from datetime import UTC, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from bot.db.enums import RaffleStatus
from bot.db.repositories import audit as audit_repo
from bot.db.session import session_scope
from bot.domain.prizes import max_prizes_per_ticket
from bot.services.draw import record_manual_winner, run_random_draw
from bot.services.errors import ExportError
from bot.services.exports import (
    ExportRequest,
    build_csv,
    build_export,
    build_tickets_pdf,
)
from bot.services.integrity import check_raffle
from bot.services.permissions import Permissions
from bot.services.reviews import confirm_application
from tests.helpers import make_application, make_organization, make_raffle, make_user

pytestmark = pytest.mark.integration
ADMIN = Permissions(user_id=100, is_admin=True)


def test_prize_cap():
    assert max_prizes_per_ticket(10, 3) == 4
    assert max_prizes_per_ticket(10, 100) == 1
    assert max_prizes_per_ticket(1, 1) == 1
    with pytest.raises(ValueError):
        max_prizes_per_ticket(0, 1)
    with pytest.raises(ValueError):
        max_prizes_per_ticket(1, 0)


async def _confirmed_raffle(session_factory):
    async with session_scope(session_factory) as session:
        org_a = await make_organization(session, name="Org A", chat_id=-1)
        org_b = await make_organization(session, name="Org B", chat_id=-2)
        raffle = await make_raffle(session, organizations=[org_a, org_b], prizes_count=3)
        user_a = await make_user(session, 501)
        user_b = await make_user(session, 502)
        a1 = await make_application(
            session, raffle=raffle, user=user_a, organization=org_a, amount=Decimal("9")
        )
        b1 = await make_application(
            session, raffle=raffle, user=user_b, organization=org_b, amount=Decimal("5")
        )
        a2 = await make_application(
            session, raffle=raffle, user=user_a, organization=org_b, amount=Decimal("6")
        )
    for application in (a1, b1, a2):
        async with session_scope(session_factory) as session:
            await confirm_application(
                session,
                application_id=application.id,
                reviewer_id=100,
                reviewer=ADMIN,
                now=application.submitted_at,
            )
    return raffle, org_a, org_b


async def test_csv_and_xlsx_contents(session_factory, env_settings):
    raffle, org_a, _ = await _confirmed_raffle(session_factory)
    tz = env_settings.timezone
    async with session_factory() as session:
        raffle = await session.get(type(raffle), raffle.id)
        csv_data = build_csv(await _load(session, raffle), raffle, tz)
        export = await build_export(
            session,
            raffle=raffle,
            format_name="xlsx",
            request=ExportRequest(),
            tz=tz,
        )
        restricted = await build_export(
            session,
            raffle=raffle,
            format_name="xlsx",
            request=ExportRequest(allowed_organization_ids=frozenset({org_a.id}), is_org_rep=True),
            tz=tz,
        )
    assert csv_data.startswith(b"\xef\xbb\xbf")
    csv_text = csv_data.decode("utf-8-sig")
    assert "application_id;payment_code;" in csv_text
    assert "9.00" in csv_text
    assert "Org A" in csv_text

    workbook = load_workbook(io.BytesIO(export.data))
    assert workbook.sheetnames == ["Applications", "Participants", "Tickets", "Summary"]
    applications = workbook["Applications"]
    assert applications.max_row == 4
    headers = [cell.value for cell in applications[1]]
    assert headers[:4] == ["raffle", "application_id", "payment_code", "organization"]
    participants = workbook["Participants"]
    assert participants.max_row == 3
    tickets = workbook["Tickets"]
    assert tickets.max_row == 5
    ticket_numbers = [tickets.cell(row=index, column=1).value for index in range(2, 6)]
    assert ticket_numbers == [1, 2, 3, 4]

    restricted_book = load_workbook(io.BytesIO(restricted.data))
    assert "Participants" not in restricted_book.sheetnames
    restricted_apps = restricted_book["Applications"]
    assert restricted_apps.max_row == 2
    assert restricted_apps.cell(row=2, column=4).value == "Org A"


async def _load(session, raffle):
    from bot.services.exports import _load_data

    return await _load_data(session, raffle, ExportRequest())


async def test_org_rep_scope_in_xlsx(session_factory, env_settings):
    raffle, org_a, _ = await _confirmed_raffle(session_factory)
    async with session_factory() as session:
        raffle = await session.get(type(raffle), raffle.id)
        export = await build_export(
            session,
            raffle=raffle,
            format_name="xlsx",
            request=ExportRequest(allowed_organization_ids=frozenset({org_a.id}), is_org_rep=True),
            tz=env_settings.timezone,
        )
    workbook = load_workbook(io.BytesIO(export.data))
    assert workbook["Applications"].max_row == 2
    assert "Participants" not in workbook.sheetnames


async def test_pdf_tickets_only_after_finish_and_count(session_factory, env_settings):
    raffle, _, _ = await _confirmed_raffle(session_factory)
    async with session_scope(session_factory) as session:
        raffle = await session.get(type(raffle), raffle.id)
        with pytest.raises(ExportError):
            build_tickets_pdf(raffle)
        raffle.status = RaffleStatus.FINISHED
        raffle.last_ticket_number = 25
    async with session_factory() as session:
        raffle = await session.get(type(raffle), raffle.id)
        export = build_tickets_pdf(raffle)
    reader = PdfReader(io.BytesIO(export.data))
    assert len(reader.pages) == 3
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "25" in text
    assert "501" not in text
    assert "Org A" not in text


async def test_random_draw_cap_and_exact_count(session_factory):
    async with session_scope(session_factory) as session:
        org = await make_organization(session)
        raffle = await make_raffle(
            session, organizations=[org], prizes_count=5, status=RaffleStatus.FINISHED
        )
        raffle.last_ticket_number = 3
        user = await make_user(session, 500)
        app = await make_application(
            session, raffle=raffle, user=user, organization=org, amount=Decimal("15")
        )
    async with session_scope(session_factory) as session:
        raffle = await session.get(type(raffle), raffle.id)
        app = await session.get(type(app), app.id)
        app.status = "confirmed"
        app.tickets_awarded = 3
        app.first_ticket_number = 1
        app.last_ticket_number = 3
        app.tickets_total_after = 3
    async with session_scope(session_factory) as session:
        raffle = await session.get(type(raffle), raffle.id)
        winners = await run_random_draw(
            session,
            actor=ADMIN,
            raffle=raffle,
            now=datetime.now(UTC),
            rng=random.Random(7),
        )
        assert len(winners) == 5
        counts: dict[int, int] = {}
        for winner in winners:
            counts[winner.ticket_number] = counts.get(winner.ticket_number, 0) + 1
        assert max(counts.values()) <= 2
        assert set(counts) <= {1, 2, 3}


async def test_manual_draw_ignores_cap_and_allows_repeat(session_factory):
    async with session_scope(session_factory) as session:
        org = await make_organization(session)
        raffle = await make_raffle(
            session, organizations=[org], prizes_count=3, status=RaffleStatus.FINISHED
        )
        raffle.last_ticket_number = 1
        user = await make_user(session, 500)
        app = await make_application(
            session, raffle=raffle, user=user, organization=org, amount=Decimal("5")
        )
    async with session_scope(session_factory) as session:
        app = await session.get(type(app), app.id)
        app.status = "confirmed"
        app.tickets_awarded = 1
        app.first_ticket_number = 1
        app.last_ticket_number = 1
        app.tickets_total_after = 1
    async with session_scope(session_factory) as session:
        raffle = await session.get(type(raffle), raffle.id)
        first = await record_manual_winner(
            session, actor=ADMIN, raffle=raffle, place=1, ticket_number=1, now=datetime.now(UTC)
        )
        second = await record_manual_winner(
            session, actor=ADMIN, raffle=raffle, place=2, ticket_number=1, now=datetime.now(UTC)
        )
        assert first.ticket_number == second.ticket_number == 1
        entries = await audit_repo.list_entries(session, action="winner_recorded")
        assert len(entries) == 2


async def test_integrity_detects_counter_mismatch(session_factory):
    raffle, _, _ = await _confirmed_raffle(session_factory)
    async with session_scope(session_factory) as session:
        raffle = await session.get(type(raffle), raffle.id)
        report = await check_raffle(session, raffle)
        assert report.is_ok
        raffle.last_ticket_number += 1
        await session.flush()
        broken = await check_raffle(session, raffle)
        assert not broken.is_ok
        codes = {issue.code for issue in broken.issues}
        assert "counter_vs_applications" in codes
        assert "counter_vs_participants" in codes
